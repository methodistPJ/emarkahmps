"""Read-only XLSX acceptance checks. Requires openpyxl; does not author/edit workbook."""
from pathlib import Path
import openpyxl,json,zipfile,hashlib,xml.etree.ElementTree as ET
root=Path(__file__).resolve().parent.parent
p=root/'assets/MPS-eAssessment-School-Setup-2026.xlsx';before=hashlib.sha256(p.read_bytes()).hexdigest();c=json.loads((root/'tools/setup-contract.json').read_text());w=openpyxl.load_workbook(p,data_only=False)
assert w.sheetnames==c['requiredSheets']
for name,fields in c['sheets'].items():
 s=w[name];assert [s.cell(4,i+1).value for i in range(len(fields))]==fields
 assert s.protection.sheet and s['A4'].protection.locked and not s['A5'].protection.locked
 assert not any(s.cell(r,i+1).data_type=='f' for r in range(5,s.max_row+1) for i in range(len(fields)))
 assert not any(f in fields for f in ['pin','proof','activation_code','AUTH_PEPPER','script_id','spreadsheet_id'])
assert w['04_TEACHERS']['C5'].value=='025849' and w['04_TEACHERS']['C5'].data_type=='s'
assert w['04_TEACHERS']['C5'].number_format=='@' and w['03_STUDENTS']['A5'].number_format=='@'
assert w['07_SUBJECTS']['E9'].value=='TB'
assert len(w['06_SUBJECT_ASSIGNMENTS'].data_validations.dataValidation)>=3
for name in ['00_START_HERE','09_VALIDATION_REFERENCE']:assert w[name].protection.sheet
with zipfile.ZipFile(p) as z:
 for n in z.namelist():
  if n.endswith('.xml') or n.endswith('.rels'):ET.fromstring(z.read(n))
assert hashlib.sha256(p.read_bytes()).hexdigest()==before
print('PASS shipped XLSX:10 worksheets, exact headers, typed text025849/IDs, data validations, protected headers/reference, editable data, no secret columns/formulas and valid XML; input hash unchanged.')
